import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from email_extractor.export import export_xlsx
from email_extractor.models import EmailRecord


class ExportTests(unittest.TestCase):
    def test_excel_contains_subject_and_not_folder(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "emails.xlsx"
            record = EmailRecord(
                datetime(2026, 9, 16, 13, 0),
                "7001234567",
                "[EXTERNAL] SALA",
                "Folha de dados",
                "Boa tarde!",
                "body original",
                "Pedido 7001234567",
            )
            self.assertEqual(export_xlsx([record], output, subject_filter="SALA"), 1)
            workbook = load_workbook(output, read_only=True)
            headers = list(workbook.active.iter_rows(min_row=1, max_row=1, values_only=True))[0]
            workbook.close()
            self.assertEqual(headers, ("Data e hora de recebimento", "Assunto", "ID", "Tipo", "Mensagem"))
            self.assertNotIn("Pasta", headers)

    def test_generic_layout_exports_untreated_body_for_blank_subject(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "emails.xlsx"
            raw_body = "<html>  Corpo original\ncom espaços  </html>"
            record = EmailRecord(
                datetime(2026, 9, 16, 13, 0),
                "7001234567",
                "Outro assunto",
                "valor tratado que não deve ser exportado",
                "mensagem tratada que não deve ser exportada",
                raw_body,
                "Inbox",
            )
            self.assertEqual(export_xlsx([record], output, subject_filter=""), 1)
            workbook = load_workbook(output, read_only=True)
            rows = list(workbook.active.iter_rows(min_row=1, max_row=2, values_only=True))
            workbook.close()
            self.assertEqual(rows[0], ("Data e hora de recebimento", "Assunto", "ID", "Body"))
            self.assertEqual(rows[1][3], raw_body)

    def test_custom_subject_uses_generic_layout(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "emails.xlsx"
            record = EmailRecord(
                datetime(2026, 9, 16, 13, 0), "", "Aviso", "tipo", "mensagem", "body", "Inbox"
            )
            export_xlsx([record], output, subject_filter="Aviso")
            workbook = load_workbook(output, read_only=True)
            headers = list(workbook.active.iter_rows(min_row=1, max_row=1, values_only=True))[0]
            workbook.close()
            self.assertEqual(headers, ("Data e hora de recebimento", "Assunto", "ID", "Body"))


if __name__ == "__main__":
    unittest.main()
